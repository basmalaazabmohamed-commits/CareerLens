
from pathlib import Path


import pymupdf
from docx import Document


def extract_cv_text(file_path: str | Path) -> str:
    """Extract text from a PDF or DOCX CV."""
    path = Path(file_path)

    if not path.is_file():
        raise FileNotFoundError(f"CV not found: {path}")

    suffix = path.suffix.lower()

    if suffix == ".pdf":
        parts = []

        with pymupdf.open(path) as pdf:
            for page_number, page in enumerate(pdf, start=1):
                page_text = page.get_text("text")

                if page_text.strip():
                    parts.append(
                        f"[Page {page_number}]\n{page_text}"
                    )

        text = "\n\n".join(parts)

    elif suffix == ".docx":
        document = Document(path)
        parts = [
            paragraph.text
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]

        # Include table content, which can contain CV details.
        for table in document.tables:
            for row in table.rows:
                values = [
                    cell.text.strip()
                    for cell in row.cells
                    if cell.text.strip()
                ]
                if values:
                    parts.append(" | ".join(values))

        text = "\n".join(parts)

    else:
        raise ValueError("Supported formats: PDF and DOCX.")

    if not text.strip():
        raise ValueError(
            "No text was extracted. The PDF may require OCR."
        )

    return text
